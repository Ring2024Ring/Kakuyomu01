import json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ids = {
    '2912051605503235345': '神運営',
    '2912051601660958663': '湖畔',
    '2912051603776752865': '奈落吸血鬼',
    '2912051605512754789': 'コンビニ',
    '2912051604996628001': 'TCG',
    '2912051597492612020': 'キャンプ',
    '2912051607205969130': '解呪師',
    '2912051602624199029': 'ブラック社員',
    '2912051606315298034': '過労死',
}
obs = {i: [] for i in ids}
files = list((ROOT / 'reports').glob('*.xlsx')) + [ROOT / 'ランキングレポート/growth_report_2026-09-30.xlsx']
for f in sorted(files, key=lambda p: p.name):
    w = openpyxl.load_workbook(f, read_only=True, data_only=True)
    if '全差分' not in w.sheetnames:
        w.close()
        continue
    it = w['全差分'].iter_rows(values_only=True)
    hs = next(it)
    for n, row in enumerate(it, 2):
        r = dict(zip(hs, row))
        wid = str(r.get('work_id'))
        if wid in ids and r.get('rank') is not None:
            obs[wid].append(dict(date=f.stem[-10:], rank=r.get('rank'), star=r.get('star'), episodes=r.get('episodes'), chars=r.get('chars'), row=n))
    w.close()
for wid, o in obs.items():
    print(json.dumps(dict(title=ids[wid], observations=len(o), first=o[0] if o else None, last=o[-1] if o else None, endpoints=[r for r in o if r['date'] in ['2026-09-18', '2026-09-24', '2026-09-30']]), ensure_ascii=False))
(Path(__file__).parent / 'growth_observations.json').write_text(json.dumps(obs, ensure_ascii=False, indent=2), encoding='utf-8')
