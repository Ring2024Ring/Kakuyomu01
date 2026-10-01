from pathlib import Path
from collections import Counter, defaultdict
import json, re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
LATEST = ROOT / 'ランキングレポート' / 'growth_report_2026-09-30.xlsx'

def records(ws):
    it = ws.iter_rows(values_only=True)
    headers = next(it)
    return [dict(zip(headers, row), source_row=i) for i, row in enumerate(it, 2) if any(v is not None for v in row)]

w = openpyxl.load_workbook(LATEST, read_only=True, data_only=True)
latest = records(w['全差分'])
current = [r for r in latest if r.get('rank') is not None]
stays = records(w['滞在日数'])
known = {}
aliases = defaultdict(set)
daily_coverage = []
files = sorted((ROOT / 'reports').glob('*.xlsx')) + [LATEST]
for path in files:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if '全差分' not in wb.sheetnames:
        wb.close()
        continue
    rows = records(wb['全差分'])
    daily_coverage.append({'file':path.name, 'genres':dict(Counter(r.get('genre') for r in rows if r.get('rank') is not None)), 'ranked':sum(r.get('rank') is not None for r in rows)})
    for row in rows:
        if row.get('work_id'):
            wid = str(row['work_id'])
            row['source_file'] = str(path.relative_to(ROOT))
            known[wid] = row
            aliases[(row.get('genre'), row.get('title'))].add(wid)
    wb.close()

current_ids = {str(r['work_id']) for r in current}
current_by_id = {str(r['work_id']):r for r in current}
history = []
for row in stays:
    ids = aliases.get((row.get('genre'),row.get('title')),set())
    match = known[next(iter(ids))] if len(ids)==1 else None
    merged = dict(match or {})
    merged.update(row)
    merged['in_current'] = str(merged.get('work_id')) in current_ids
    merged['history_source_row'] = row['source_row']
    merged['current_source_row'] = current_by_id.get(str(merged.get('work_id')),{}).get('source_row')
    history.append(merged)

history_unique = {}
for row in history:
    key = ('id',str(row['work_id'])) if row.get('work_id') else ('title',row.get('genre'),row.get('title'))
    if key not in history_unique:
        history_unique[key] = row
    elif row['history_source_row'] > history_unique[key]['history_source_row']:
        history_unique[key] = row
history_deduped = list(history_unique.values())

dev = re.compile('経営|開拓|開発|運営|管理|改築|改装|建設|拠点|農|工場|温泉|ホテル|貸|利用権|採掘|領主|領地|支配|復興|再生|修復|鍛冶|生産|商売|資源|宿|店')
monster = re.compile('吸血鬼|ヴァンパイア|バンパイア|眷属|血族|領域|ネクロ|アンデッド|死霊')
territory = re.compile('経営|領地|領主|復興|支配|配下|再建|修復|再生')
def screening(row):
    intro = row.get('intro') or row.get('あらすじ') or ''
    dungeon = bool(re.search('ダンジョン|迷宮', intro))
    flags=[]
    if dungeon and dev.search(intro): flags.append('dungeon_development_screen')
    if monster.search(intro): flags.append('monster_or_domain_screen')
    if territory.search(intro): flags.append('territory_screen')
    return flags

for name, rows in [('latest_records', latest), ('current_records', current), ('history_records', history), ('history_deduped',history_deduped), ('daily_coverage',daily_coverage)]:
    (OUT / (name+'.json')).write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding='utf-8')

screened = [dict(r, screening=screening(r)) for r in history_deduped if screening(r)]
(OUT/'screened_records.json').write_text(json.dumps(screened, ensure_ascii=False, indent=2, default=str),encoding='utf-8')
summary = {
    'latest_rows':len(latest),'current':len(current),'current_unique_ids':len(current_ids),
    'current_genres':dict(Counter(r.get('genre') for r in current)),
    'current_status':dict(Counter(r.get('status') for r in current)),
    'current_classes':dict(Counter(r.get('区分') for r in current)),
    'current_span_days':dict(Counter(r.get('span_days') for r in current)),
    'history_rows':len(history),'history_linked_ids':sum(bool(r.get('work_id')) for r in history),
    'history_deduped':len(history_deduped),'history_unique_linked_ids':sum(bool(r.get('work_id')) for r in history_deduped),
    'history_genres':dict(Counter(r.get('genre') for r in history)),
    'screened':len(screened),'screened_current':sum(r['in_current'] for r in screened),
    'current_dungeon_intro_keyword':sum(bool(re.search('ダンジョン|迷宮',r.get('intro') or r.get('あらすじ') or '')) for r in current),
    'report_files_read':len(daily_coverage),
    'latest_daily_coverage':daily_coverage[-4:],
}
(OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
print('\nHISTORICAL SCREENED CANDIDATES')
for r in screened:
    if not r['in_current']:
        print(json.dumps({k:r.get(k) for k in ['genre','title','work_id','最高順位','最終順位','滞在日数','観測回数','status','screening']},ensure_ascii=False))
