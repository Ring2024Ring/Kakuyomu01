from pathlib import Path
from collections import Counter
from html.parser import HTMLParser
import json,sys,datetime
import openpyxl
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'analysis'
class Extract(HTMLParser):
    def __init__(self):
        super().__init__(); self.active=None; self.parts=[]; self.items=[]
    def handle_starttag(self,tag,attrs):
        a=dict(attrs)
        if tag in ('h2','h3','footer') or (tag=='div' and a.get('class') in ('sub','note','danger')):
            self.active=tag;self.parts=[]
    def handle_data(self,d):
        if self.active:self.parts.append(d)
    def handle_endtag(self,tag):
        if tag==self.active:
            self.items.append(' '.join(''.join(self.parts).split()));self.active=None
allrows=[]; summaries=[]
for p in sorted((ROOT/'reports').glob('*.xlsx')):
    date=p.stem[-10:]
    w=openpyxl.load_workbook(p,read_only=True,data_only=True)
    candidates=[s for s in w if s.title=='全差分']
    if not candidates:
        print('NO FULL DATA',date,w.sheetnames);continue
    s=candidates[0];vals=iter(s.values);headers=next(vals)
    rows=[dict(zip(headers,r)) for r in vals]
    for r in rows:r['date']=date
    allrows.extend(rows)
    html=Extract();html.feed(p.with_suffix('.html').read_text(encoding='utf-8'))
    summaries.append({'date':date,'rows':len(rows),'unique':len({r['work_id'] for r in rows}),'genres':dict(Counter(r['genre'] for r in rows)),'ranks':{g:[min(r['rank'] for r in rows if r['genre']==g and r['rank'] is not None),max(r['rank'] for r in rows if r['genre']==g and r['rank'] is not None)] for g in {r['genre'] for r in rows if r['rank'] is not None}},'notes':html.items,'sheets':[(sh.title,sh.max_row,sh.max_column) for sh in w]})
    w.close()
(OUT/'observations.json').write_text(json.dumps(allrows,ensure_ascii=False,default=str),encoding='utf8')
(OUT/'report_metadata.json').write_text(json.dumps(summaries,ensure_ascii=False,indent=2),encoding='utf8')
print('TOTAL',len(summaries),'reports',len(allrows),'rows',len({r['work_id'] for r in allrows}),'unique IDs')
print('DATE_COUNTS',[(s['date'],s['rows'],s['unique'],s['genres']) for s in summaries])
print('LATEST_META',json.dumps(summaries[-1],ensure_ascii=False,indent=2))
print('HEADERS',list(allrows[-1]))
print('LATEST_TOP20')
latest=[r for r in allrows if r['date']==summaries[-1]['date']]
for r in sorted(latest,key=lambda r:-(r.get('star_per_day') or 0))[:20]:
    print(json.dumps({k:r.get(k) for k in ('work_id','genre','rank','title','star','star_gain','star_per_day','elapsed_days','span_days','区分')},ensure_ascii=False))

